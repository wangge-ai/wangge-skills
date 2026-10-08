import argparse
import os
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = Path(os.environ.get("WANGGE_VIDEO_OUTPUT", str(Path.cwd() / "outputs" / "data-video")))
SKILL_ROOT = ROOT


SOURCE_COLUMNS = [
    "category",
    "source_name",
    "owner_or_publisher",
    "official_level",
    "home_url",
    "api_or_download",
    "coverage",
    "typical_metrics",
    "hot_keywords",
    "video_formats",
    "notes",
    "last_reviewed",
    "refresh_cadence",
    "source_mode",
    "live_search_query",
    "refresh_rule",
]

ROUTE_COLUMNS = [
    "hot_topic_type",
    "trigger_keywords",
    "preferred_sources",
    "backup_sources",
    "default_video_format",
    "data_shape",
    "example_titles",
    "quality_gate",
    "source_mode",
    "live_search_required_when",
    "live_search_queries",
    "refresh_cadence",
]


def read_table(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig").fillna("")


def source_refresh_cadence(row: pd.Series) -> str:
    category = str(row.get("category", ""))
    level = str(row.get("official_level", ""))
    if any(key in category for key in ["灾害", "空气质量", "金融", "旅游", "交通"]):
        return "每次使用前校验，至少每7天复查"
    if any(key in level for key in ["第三方", "专业统计", "专业整理"]):
        return "每次使用前校验，至少每30天复查"
    return "每次使用前校验，至少每90天复查"


def source_live_query(row: pd.Series) -> str:
    parts = [
        str(row.get("source_name", "")),
        str(row.get("owner_or_publisher", "")),
        str(row.get("category", "")),
        str(row.get("typical_metrics", "")),
        "official data API CSV download documentation",
    ]
    return " ".join(part for part in parts if part).strip()


def route_live_queries(row: pd.Series) -> str:
    topic = str(row.get("hot_topic_type", ""))
    keywords = " ".join(str(row.get("trigger_keywords", "")).split("|")[:4])
    return "|".join(
        [
            f"{topic} 官方 数据 API CSV 下载",
            f"{topic} official data API csv dataset",
            f"{keywords} 数据源 官方 统计 API",
        ]
    )


def prepare_source_table(df: pd.DataFrame) -> pd.DataFrame:
    today = date.today().isoformat()
    out = df.copy()
    out["last_reviewed"] = out.get("last_reviewed", today)
    out.loc[out["last_reviewed"].astype(str).str.len() == 0, "last_reviewed"] = today
    out["refresh_cadence"] = out.apply(source_refresh_cadence, axis=1)
    out["source_mode"] = "种子库，不是最终答案；使用前必须联网校验来源是否仍可用"
    out["live_search_query"] = out.apply(source_live_query, axis=1)
    out["refresh_rule"] = "先打开 home_url/api_or_download 核对；若失效、口径不匹配或热点不覆盖，按 live_search_query 联网补源并回写新来源"
    return out[[column for column in SOURCE_COLUMNS if column in out.columns]]


def prepare_route_table(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["source_mode"] = "先查种子库；不命中、过期或口径不合适时必须联网搜索补源"
    out["live_search_required_when"] = "热点未命中本表|首选来源不可用|来源超过刷新周期|数据没有所需时间/地区/对象|只有二手截图或无法复查"
    out["live_search_queries"] = out.apply(route_live_queries, axis=1)
    out["refresh_cadence"] = "每次使用前快速校验；每月批量复查路由和首选来源"
    return out[[column for column in ROUTE_COLUMNS if column in out.columns]]


def autosize_sheet(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = True
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx in range(1, sheet.max_column + 1):
            letter = get_column_letter(col_idx)
            values = [str(sheet.cell(row=row_idx, column=col_idx).value or "") for row_idx in range(1, min(sheet.max_row, 80) + 1)]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 42)
            sheet.column_dimensions[letter].width = width
        sheet.row_dimensions[1].height = 34
    workbook.save(workbook_path)


def write_workbooks(source_df: pd.DataFrame, route_df: pd.DataFrame, outputs: Path) -> list[Path]:
    combined = outputs / "TOP10stats_数据源和路由表_Excel版.xlsx"
    source_xlsx = outputs / "官方专业数据源库.xlsx"
    route_xlsx = outputs / "热点到数据视频路由表.xlsx"

    note_rows = pd.DataFrame(
        [
            {"说明": "CSV 已改为 UTF-8 BOM，Excel 双击打开不应再乱码；更推荐直接打开本 xlsx。"},
            {"说明": "专业数据源库只是种子库，不是写死清单。每次制作前要按 refresh_cadence 和 live_search_query 联网校验。"},
            {"说明": "热点如果不在路由表或来源不合适，必须联网搜索官方/专业数据源，再把新来源回写到库里。"},
        ]
    )
    with pd.ExcelWriter(combined, engine="openpyxl") as writer:
        note_rows.to_excel(writer, sheet_name="使用说明", index=False)
        source_df.to_excel(writer, sheet_name="官方专业数据源库", index=False)
        route_df.to_excel(writer, sheet_name="热点路由表", index=False)
    with pd.ExcelWriter(source_xlsx, engine="openpyxl") as writer:
        source_df.to_excel(writer, sheet_name="官方专业数据源库", index=False)
    with pd.ExcelWriter(route_xlsx, engine="openpyxl") as writer:
        route_df.to_excel(writer, sheet_name="热点路由表", index=False)
    for path in [combined, source_xlsx, route_xlsx]:
        autosize_sheet(path)
    return [combined, source_xlsx, route_xlsx]


def sync_skill_references(source_csv: Path, route_csv: Path, outputs: Path) -> None:
    references = SKILL_ROOT / "references"
    if not references.exists():
        return
    (references / "official-data-source-library.csv").write_bytes(source_csv.read_bytes())
    (references / "hotspot-to-video-routing.csv").write_bytes(route_csv.read_bytes())
    for name in ["TOP10stats_数据源和路由表_Excel版.xlsx", "官方专业数据源库.xlsx", "热点到数据视频路由表.xlsx"]:
        src = outputs / name
        if src.exists():
            (references / name).write_bytes(src.read_bytes())


def main() -> int:
    parser = argparse.ArgumentParser(description="Refresh reference tables into Excel-safe CSV and XLSX files.")
    parser.add_argument("--outputs", default=str(OUTPUTS), help="Outputs directory containing the two reference CSV files.")
    parser.add_argument("--sync-skill", action="store_true", help="Copy refreshed tables into the installed skill references.")
    args = parser.parse_args()

    outputs = Path(args.outputs).resolve()
    source_csv = outputs / "官方专业数据源库.csv"
    route_csv = outputs / "热点到数据视频路由表.csv"

    source_df = prepare_source_table(read_table(source_csv))
    route_df = prepare_route_table(read_table(route_csv))

    source_df.to_csv(source_csv, index=False, encoding="utf-8-sig")
    route_df.to_csv(route_csv, index=False, encoding="utf-8-sig")
    workbooks = write_workbooks(source_df, route_df, outputs)
    if args.sync_skill:
        sync_skill_references(source_csv, route_csv, outputs)

    print("refreshed:")
    print(source_csv)
    print(route_csv)
    for path in workbooks:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
