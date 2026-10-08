#!/usr/bin/env python3
import argparse
import json
import os
import subprocess
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

parser = argparse.ArgumentParser()
parser.add_argument('--task-id', required=True)
parser.add_argument('--artifact-id', required=True)
parser.add_argument('--input-manifest', required=True)
parser.add_argument('--output-dir', required=True)
args = parser.parse_args()

TASK_CONTRACTS = {
    'video-analysis': {'primary_output': 'video-evidence-brief', 'workflow_sections': ['hook', 'product-fact', 'visible-proof', 'call-to-action'], 'acceptance_focus': ['trace each judgement to local video or structured evidence', 'hold unconfirmed selling points']},
    'detail-page-planning': {'primary_output': 'detail-page-screen-plan', 'workflow_sections': ['first-screen', 'pain-point', 'solution-proof', 'details', 'trust', 'call-to-action'], 'acceptance_focus': ['map every screen to evidence or a material gap', 'do not state unverified parameters as fact']},
    'main-image-series-planning': {'primary_output': 'main-image-role-plan', 'workflow_sections': ['hero', 'core-benefit', 'scene', 'detail-proof', 'spec-boundary', 'trust'], 'acceptance_focus': ['give each image one communication role', 'do not create or publish final images']},
    'main-image-marketing-positioning': {'primary_output': 'main-image-position-map', 'workflow_sections': ['audience', 'purchase-scene', 'pain-point', 'credible-benefit', 'hero-expression'], 'acceptance_focus': ['separate positioning claims from supplied evidence', 'do not claim competitive advantage without market evidence']},
    'main-image-selling-point-analysis': {'primary_output': 'selling-point-evidence-map', 'workflow_sections': ['visible-benefit', 'proof', 'expression-risk', 'material-gap', 'priority'], 'acceptance_focus': ['map every selling point to evidence and risk', 'do not reuse competitor copy or unverified certifications']},
    'video-planning': {'primary_output': 'video-shot-list', 'workflow_sections': ['opening', 'shots', 'voiceover', 'shooting-assets', 'acceptance'], 'acceptance_focus': ['use only confirmed product facts', 'do not generate final video']},
    'marketing-video-planning': {'primary_output': 'marketing-video-content-matrix', 'workflow_sections': ['content-theme', 'audience-scene', 'selling-point-proof', 'script-direction', 'channel-rhythm'], 'acceptance_focus': ['label evidence status for every recommendation', 'do not call external video generation or publishing services']}
}
task_contract = TASK_CONTRACTS.get(args.task_id, {
    'primary_output': 'evidence-bounded-planning-artifact',
    'workflow_sections': ['fact-check', 'work-plan', 'acceptance-boundary'],
    'acceptance_focus': ['trace conclusions to user-provided material']
})

manifest = json.loads(Path(args.input_manifest).read_text(encoding='utf-8'))
out = Path(args.output_dir)
out.mkdir(parents=True, exist_ok=True)
inputs = manifest.get('inputs', {})
evidence = list(inputs.get('evidence', []))
videos = list(inputs.get('videos', []))
images = list(inputs.get('images', []))
video_meta = []
for video in videos:
    item = {'path': video, 'bytes': os.path.getsize(video), 'metadata_status': 'unavailable'}
    try:
        probe = subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'json', video], capture_output=True, text=True, timeout=15, check=False)
        if probe.returncode == 0:
            item.update({'metadata_status': 'ffprobe', 'duration_seconds': json.loads(probe.stdout).get('format', {}).get('duration')})
    except (FileNotFoundError, subprocess.TimeoutExpired, json.JSONDecodeError):
        pass
    video_meta.append(item)

facts = {
    'task_id': args.task_id,
    'artifact_id': args.artifact_id,
    'status': 'partial',
    'evidence_count': len(evidence) + len(videos) + len(images),
    'evidence_files': evidence,
    'video_files': video_meta,
    'task_contract': task_contract,
    'image_files': images,
    'held_claims': ['所有未出现在用户提供证据中的参数、认证、效果、价格和促销信息均待核实'],
    'generated_media': 0,
    'external_provider_called': False
}
(out / f'{args.artifact_id}.json').write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding='utf-8')
(out / f'{args.artifact_id}.md').write_text(
    f'# {args.artifact_id}\n\n- 业务状态：partial\n- 证据文件：{facts["evidence_count"]}\n- 成品媒体：0\n- 外部生成服务：未调用\n\n## 待核实\n\n- {facts["held_claims"][0]}\n', encoding='utf-8')
with (out / f'{args.artifact_id}.md').open('a', encoding='utf-8') as markdown:
    markdown.write('\n## Task contract\n')
    markdown.write(f'- Primary output: {task_contract["primary_output"]}\n')
    markdown.write(f'- Workflow: {" / ".join(task_contract["workflow_sections"])}\n')
    markdown.write(f'- Acceptance: {" / ".join(task_contract["acceptance_focus"])}\n')

wb = Workbook()
overview = wb.active
overview.title = '概览'
overview.append(['字段', '内容'])
overview.append(['任务', args.task_id])
overview.append(['证据状态', '用户提供；未验证项保留'])
overview.append(['成品媒体', 0])
overview.append(['外部生成服务', '未调用'])
evidence_sheet = wb.create_sheet('证据与边界')
evidence_sheet.append(['文件', '类型', '状态'])
for file in evidence: evidence_sheet.append([file, '结构化证据', '用户提供'])
for file in images: evidence_sheet.append([file, '图片', '用户提供'])
for item in video_meta: evidence_sheet.append([item['path'], '视频', item['metadata_status']])
tasks = wb.create_sheet('执行任务')
tasks.append(['步骤', '任务', '证据状态', '验收'])
tasks.append([1, '核对事实与素材', '用户提供', '每项结论可追溯到输入文件'])
tasks.append([2, '形成策划或分析结论', '部分完成', '未核实内容保持待核实'])
tasks.append([3, '外部生成或发布', '未执行', '不属于本任务范围'])
design = wb.create_sheet('Task contract')
design.append(['Item', 'Content'])
design.append(['Primary output', task_contract['primary_output']])
design.append(['Workflow', ' / '.join(task_contract['workflow_sections'])])
design.append(['Acceptance', ' / '.join(task_contract['acceptance_focus'])])
for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    sheet.column_dimensions['A'].width = 28
    sheet.column_dimensions['B'].width = 60
    sheet.column_dimensions['C'].width = 24
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(wrap_text=True)
wb.save(out / f'{args.artifact_id}.xlsx')
print(out / f'{args.artifact_id}.json')
